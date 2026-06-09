import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def main():
    doc_path = r"C:\Users\Marcus Schneider\Documents\Marcus\2026\Projetos\Fase_final_Apresentação\ALEP\Projeto ALEP_Finalização pendente.docx"
    out_path = r"C:\Users\Marcus Schneider\Documents\Marcus\2026\Projetos\Fase_final_Apresentação\ALEP\Projeto_ALEP_Finalizado.docx"
    
    doc = docx.Document(doc_path)
    
    # 6.2. A Ferramenta como Indutora de Maturidade Política e Elucidação Cidadã
    p_h2 = doc.add_paragraph()
    p_h2.style = 'Heading 2' if 'Heading 2' in doc.styles else 'Normal'
    run_h2 = p_h2.add_run("6.2. A Ferramenta como Indutora de Maturidade Política e Elucidação Cidadã")
    run_h2.bold = True
    
    p1 = doc.add_paragraph(
        "A implementação de tecnologias cívicas (Civic Techs) transcende a automação de processos "
        "administrativos; ela atua como um vetor pedagógico no tecido social. Ao converter o complexo Big Data governamental "
        "em uma interface analítica acessível, a ferramenta desenvolvida exerce um papel fundamental na "
        "elucidação do cidadão. Conforme demonstrado, a verdadeira transparência não reside apenas na disponibilidade do dado bruto, "
        "mas na capacidade cognitiva do receptor em compreendê-lo."
    )
    p1.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    
    p2 = doc.add_paragraph(
        "Nesta esteira, o painel interativo (dashboard) estabelece um novo padrão para o controle social extraorgânico. "
        "Ao fragmentar e mapear a atuação legislativa — distinguindo claramente o que é inovação do ordenamento jurídico (Projetos de Lei "
        "e Normas Legais) daquelas ações voltadas à articulação regional stricto sensu (Requerimentos, Indicações e Votos de Louvor) —, "
        "o cidadão é empoderado com a métrica real do mandato de seus representantes. Rompe-se, dessa forma, a dependência nociva "
        "dos recortes otimizados por algoritmos de marketing político nas redes sociais, que muitas vezes distorcem a percepção pública "
        "ao invisibilizar as etapas menos glamurosas, porém vitais, do processo legislativo."
    )
    p2.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    
    p3 = doc.add_paragraph(
        "Isso induz, invariavelmente, ao amadurecimento político. O eleitor deixa a condição de mero espectador "
        "passivo de uma narrativa unilateral para assumir a postura de auditor cívico. Munido de dados segmentados de maneira imparcial "
        "e determinística, o cidadão torna-se capaz de correlacionar as promessas de campanha com a produção material fática do parlamentar, "
        "elevar a qualidade do debate público e realizar um escrutínio representativo muito mais técnico, consciente e qualificado nas esferas democráticas do Estado."
    )
    p3.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    
    # 7. CONCLUSÃO
    p_h1 = doc.add_paragraph()
    p_h1.style = 'Heading 1' if 'Heading 1' in doc.styles else 'Normal'
    run_h1 = p_h1.add_run("7. CONCLUSÃO")
    run_h1.bold = True
    
    c1 = doc.add_paragraph(
        "O presente projeto acadêmico e tecnológico propôs-se a enfrentar uma das maiores disfunções práticas do Estado "
        "Democrático de Direito na era digital: a dicotomia entre a publicidade formal imposta pela norma e a barreira da opacidade cognitiva. "
        "A premissa inicial de que a “obscuridade por volume” inviabilizava o controle social sobre o Poder Legislativo estadual comprovou-se "
        "fática ao se analisar o imenso acervo de matrizes JSON não consolidadas que os cidadãos precisarão interpretar isoladamente "
        "caso não contem com interfaces intermédias."
    )
    c1.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    
    c2 = doc.add_paragraph(
        "A pesquisa entregou, portanto, não apenas uma tese jurídica de aplicabilidade do Princípio da Publicidade, mas a "
        "materialidade de um ecossistema de software operável: a arquitetura E.N.E.I.A.S. Atrelada intrinsecamente aos mandamentos de *Privacy by Design* "
        "(Governança Efêmera via diretórios estritamente de trânsito) estabelecidos pela Lei Geral de Proteção de Dados (LGPD) e à "
        "preocupação com vieses algorítmicos levantada pelo Marco Legal da IA (PL 2338/2023), a solução adotou uma postura conservadora "
        "e segura. Ao dispensar o uso de redes neurais subjetivas na interpretação textual, a arquitetura assegurou 100% de transparência e determinismo, "
        "classificando as proposições exclusivamente através de filtragem semântica estruturada pelas chaves oficiais da Assembleia "
        "Legislativa do Estado do Paraná (ALEP)."
    )
    c2.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    
    c3 = doc.add_paragraph(
        "A conversão final desse complexo arranjo tecnológico na entrega de um Static Payload (arquivos JSON pré-digeridos injetados em um "
        "aplicativo Front-End interativo sem dependência de banco de dados nativo) materializa a sustentabilidade econômica e a vocação perene das *Civic Techs*. "
        "O painel analítico gerado atende plenamente aos objetivos traçados, fragmentando e evidenciando a pluralidade de atuações inerentes ao mandato legislativo, "
        "sem emitir valor de mérito."
    )
    c3.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    
    c4 = doc.add_paragraph(
        "Conclui-se, destarte, que o emprego estratégico da Engenharia de Software e da Jurimetria atua como a infraestrutura garantista "
        "do princípio republicano e da publicidade prevista no caput do Art. 37 da Constituição Federal. Ao transmutar dados densos em informação acessível, "
        "a ferramenta desmancha oligopólios narrativos da era das redes sociais e devolve o escrutínio orgânico e fidedigno a quem, por direito originário, "
        "pertence todo o poder: a sociedade civil representativa."
    )
    c4.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    
    doc.save(out_path)
    print(f"Arquivo gerado com sucesso em: {out_path}")

if __name__ == "__main__":
    main()
